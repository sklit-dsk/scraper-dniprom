from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import time
import html as html_lib
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import urljoin
from xml.etree import ElementTree as ET

import requests
from bs4 import BeautifulSoup
from bs4.element import Tag
from openai import OpenAI
from openpyxl.styles import Alignment
from openpyxl import Workbook, load_workbook


BASE_URL = "https://dnipro-m.ua"
SITEMAP_INDEX_URL = f"{BASE_URL}/sitemap_index.xml"
SITEMAP_CACHE = Path(".cache") / "sitemap_urls.json"
INDEX_CACHE = Path(".cache") / "article_url_index.json"
TRANSLATION_CACHE = Path(".cache") / "translation_cache.json"
DEBUG_LOG = Path(".cache") / "scraper_debug.log"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/126.0.0.0 Safari/537.36"
    )
}


@dataclass
class ProductRecord:
    article_input: str
    manufacturer_code: str = ""
    product_name: str = ""
    product_description: str = ""
    characteristics: str = ""
    image_1: str = ""
    image_2: str = ""
    image_3: str = ""
    source_url: str = ""
    status: str = "ok"
    error: str = ""


def ensure_cache_dir() -> None:
    Path(".cache").mkdir(parents=True, exist_ok=True)


def log_event(level: str, message: str, **context: object) -> None:
    ensure_cache_dir()
    payload: dict[str, object] = {
        "ts": time.strftime("%Y-%m-%d %H:%M:%S"),
        "level": level,
        "message": message,
    }
    payload.update(context)
    try:
        with DEBUG_LOG.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
    except Exception:
        pass


def load_text_codes(path: Path) -> list[str]:
    codes: list[str] = []
    with path.open("r", encoding="utf-8-sig") as handle:
        for line in handle:
            code = line.strip().split(",")[0].split(";")[0].strip()
            if code:
                codes.append(code)
    return codes


def load_tabular_codes(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        codes: list[str] = []
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row:
                continue
            value = row[0]
            if value is None:
                continue
            code = str(value).strip()
            if code:
                codes.append(code)
        return codes
    if suffix in {".csv"}:
        codes = []
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            for row in reader:
                if not row:
                    continue
                code = row[0].strip()
                if code:
                    codes.append(code)
        return codes
    return load_text_codes(path)


def load_article_codes(input_path: Path | None, direct_codes: list[str]) -> list[str]:
    codes: list[str] = []

    for value in direct_codes:
        code = normalize_code(value)
        if code:
            codes.append(code)

    if input_path is not None:
        codes.extend(load_tabular_codes(input_path))

    normalized = [code for code in (normalize_code(item) for item in codes) if code]
    deduplicated = list(dict.fromkeys(normalized))
    return deduplicated


def normalize_code(value: str) -> str:
    code = re.sub(r"\s+", "", value).strip()
    if re.fullmatch(r"\d{7}", code):
        return f"0{code}"
    return code


def canonical_lookup_code(value: str) -> str:
    code = normalize_code(value)
    if not code:
        return ""
    # Some SKUs on the site include suffixes like "-1"; lookup should use the base numeric part.
    match = re.match(r"^(\d{7,})", code)
    if match:
        return normalize_code(match.group(1))
    return code


def extract_product_code_from_html(html: str) -> str:
    patterns = [
        r'\bcode="(\d+)"',
        r'\bproduct-sku="(\d+)"',
        r'\bproductKey&quot;:&quot;(\d+)&quot;',
        r'"code"\s*:\s*"?(\d{7,})"?',
        r'"sku"\s*:\s*"?(\d{7,})"?',
        r'"article"\s*:\s*"?(\d{7,})"?',
        r'"productSku"\s*:\s*"?(\d{7,})"?',
    ]
    for pattern in patterns:
        match = re.search(pattern, html)
        if match:
            return normalize_code(match.group(1))
    return ""


def fetch_html(session: requests.Session, url: str, retries: int = 2, retry_delay: float = 0.6) -> str:
    retry_statuses = {400, 408, 425, 429, 500, 502, 503, 504}
    last_error: Exception | None = None
    for attempt in range(retries + 1):
        try:
            response = session.get(url, headers=HEADERS, timeout=45)
            if response.status_code in retry_statuses and attempt < retries:
                log_event(
                    "warn",
                    "retryable_http_status",
                    url=url,
                    status=response.status_code,
                    attempt=attempt + 1,
                    retries=retries,
                )
                time.sleep(retry_delay * (attempt + 1))
                continue
            response.raise_for_status()
            return response.text
        except requests.RequestException as exc:
            last_error = exc
            status_code = exc.response.status_code if getattr(exc, "response", None) is not None else None
            log_event(
                "warn",
                "http_request_exception",
                url=url,
                status=status_code,
                attempt=attempt + 1,
                retries=retries,
                error=str(exc),
            )
            if attempt >= retries:
                log_event("error", "http_request_failed", url=url, status=status_code, error=str(exc))
                raise
            time.sleep(retry_delay * (attempt + 1))

    if last_error:
        raise last_error
    raise RuntimeError(f"Failed to fetch URL: {url}")


def fetch_html_with_recovery(session: requests.Session, url: str) -> str:
    try:
        return fetch_html(session, url)
    except requests.HTTPError as exc:
        status_code = exc.response.status_code if exc.response is not None else None
        if status_code != 400:
            raise

        log_event("warn", "retry_with_fresh_session", url=url, status=status_code)
        fresh_session = requests.Session()
        try:
            return fetch_html(fresh_session, url, retries=1)
        finally:
            fresh_session.close()


def extract_product_links_from_html(html: str) -> list[str]:
    links = re.findall(r'https://dnipro-m\.ua/(?:ru|ua)/tovar/[^"\'\s<]+', html)
    deduped: list[str] = []
    seen: set[str] = set()
    for url in links:
        if url not in seen:
            seen.add(url)
            deduped.append(url)
    return deduped


def get_csrf_token(session: requests.Session) -> str:
    response = session.get(f"{BASE_URL}/ru/", headers={"User-Agent": HEADERS["User-Agent"]}, timeout=45)
    response.raise_for_status()
    match = re.search(r'<meta name="csrf-token" content="([^"]+)"', response.text)
    if not match:
        return ""
    return match.group(1)


def search_product_via_search_page(session: requests.Session, code: str, use_fresh_session: bool = False) -> str:
    normalized_code = canonical_lookup_code(code)
    if not normalized_code:
        return ""

    try:
        active_session = requests.Session() if use_fresh_session else session
        csrf_token = get_csrf_token(active_session)
    except Exception:
        csrf_token = ""

    headers = {
        "User-Agent": HEADERS["User-Agent"],
        "Accept": "application/json, text/plain, */*",
        "X-Requested-With": "XMLHttpRequest",
        "Referer": f"{BASE_URL}/ru/catalog/search/?text={normalized_code}",
    }
    if csrf_token:
        headers["X-CSRF-Token"] = csrf_token

    try:
        response = active_session.get(f"{BASE_URL}/ru/search/page/?q={normalized_code}", headers=headers, timeout=45)
        response.raise_for_status()
        data = response.json()
    except Exception:
        return ""

    products: list[dict[str, object]] = []
    if isinstance(data, dict):
        raw_products = data.get("products", {})
        if isinstance(raw_products, dict):
            candidate_products = raw_products.get("products", [])
            if isinstance(candidate_products, list):
                products = [item for item in candidate_products if isinstance(item, dict)]
        elif isinstance(raw_products, list):
            products = [item for item in raw_products if isinstance(item, dict)]
    elif isinstance(data, list):
        products = [item for item in data if isinstance(item, dict)]

    for item in products:
        if not isinstance(item, dict):
            continue
        product_code = canonical_lookup_code(str(item.get("code", "")))
        if product_code != normalized_code:
            continue
        link = str(item.get("link", "")).strip()
        if not link:
            continue
        if link.startswith("/"):
            return urljoin(BASE_URL, link)
        return link

    return ""


def discover_product_url(session: requests.Session, code: str, product_urls: list[str]) -> str:
    normalized_code = canonical_lookup_code(code)
    if not normalized_code:
        return ""

    search_page_url = search_product_via_search_page(session, normalized_code)
    if search_page_url:
        return search_page_url

    search_page_url = search_product_via_search_page(session, normalized_code, use_fresh_session=True)
    if search_page_url:
        log_event("info", "search_page_recovered_with_fresh_session", code=normalized_code, url=search_page_url)
        return search_page_url

    direct_candidates = [
        f"{BASE_URL}/ru/tovar/{normalized_code}/",
        f"{BASE_URL}/ua/tovar/{normalized_code}/",
        f"{BASE_URL}/tovar/{normalized_code}/",
    ]
    for url in direct_candidates:
        try:
            html = fetch_html(session, url, retries=1)
        except Exception:
            continue
        extracted_code = extract_product_code_from_html(html)
        if extracted_code == normalized_code:
            return url

    search_candidates = [
        f"{BASE_URL}/ru/search/?q={normalized_code}",
        f"{BASE_URL}/ua/search/?q={normalized_code}",
        f"{BASE_URL}/search/?q={normalized_code}",
    ]
    for search_url in search_candidates:
        try:
            response = session.get(search_url, headers=HEADERS, timeout=45)
            html = response.text
        except Exception:
            continue

        for url in extract_product_links_from_html(html):
            try:
                product_html = fetch_html(session, url, retries=1)
            except Exception:
                continue
            extracted_code = extract_product_code_from_html(product_html)
            if extracted_code == normalized_code:
                return url

    token_matches = [url for url in product_urls if normalized_code in url]
    for url in token_matches[:10]:
        try:
            html = fetch_html(session, url, retries=1)
        except Exception:
            continue
        extracted_code = extract_product_code_from_html(html)
        if extracted_code == normalized_code:
            return url

    return ""


def parse_sitemap_urls(xml_text: str) -> list[str]:
    root = ET.fromstring(xml_text)
    ns = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    urls = []
    for loc in root.findall(".//sm:url/sm:loc", ns):
        if loc.text and "/tovar/" in loc.text:
            urls.append(loc.text.strip())
    return urls


def get_product_urls(session: requests.Session) -> list[str]:
    ensure_cache_dir()
    if SITEMAP_CACHE.exists():
        return json.loads(SITEMAP_CACHE.read_text(encoding="utf-8"))

    index_xml = fetch_html(session, SITEMAP_INDEX_URL)
    root = ET.fromstring(index_xml)
    ns = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    sitemap_urls = [loc.text.strip() for loc in root.findall(".//sm:sitemap/sm:loc", ns) if loc.text]

    product_urls: list[str] = []
    for sitemap_url in sitemap_urls:
        xml_text = fetch_html(session, sitemap_url)
        product_urls.extend(parse_sitemap_urls(xml_text))

    product_urls = sorted(set(product_urls))
    SITEMAP_CACHE.write_text(json.dumps(product_urls, ensure_ascii=False, indent=2), encoding="utf-8")
    return product_urls


def load_article_index() -> dict[str, str]:
    if not INDEX_CACHE.exists():
        return {}
    return json.loads(INDEX_CACHE.read_text(encoding="utf-8"))


def save_article_index(index: dict[str, str]) -> None:
    ensure_cache_dir()
    INDEX_CACHE.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")


def load_translation_cache() -> dict[str, dict[str, str]]:
    if not TRANSLATION_CACHE.exists():
        return {}
    return json.loads(TRANSLATION_CACHE.read_text(encoding="utf-8"))


def save_translation_cache(cache: dict[str, dict[str, str]]) -> None:
    ensure_cache_dir()
    TRANSLATION_CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def make_translation_key(product_name: str, product_description: str, characteristics: str) -> str:
    payload = json.dumps(
        {
            "product_name": product_name,
            "product_description": product_description,
            "characteristics": characteristics,
        },
        ensure_ascii=False,
        sort_keys=True,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def extract_text_after_label(text: str, label: str) -> str:
    pattern = re.compile(re.escape(label) + r"\s*[:\-]?\s*(.+)", re.IGNORECASE)
    match = pattern.search(text)
    if match:
        return match.group(1).strip()
    return ""


def _extract_prop_value(html_text: str, prop_name: str) -> str:
    pattern = re.compile(rf'{re.escape(prop_name)}=(?P<quote>["\'])((?:.|\n)*?)(?P=quote)')
    match = pattern.search(html_text)
    if not match:
        return ""
    value = html_lib.unescape(match.group(2))
    return value.strip()


def extract_characteristics_from_raw_html(html_text: str) -> str:
    raw = _extract_prop_value(html_text, ':characteristics')
    if raw:
        try:
            data = json.loads(raw)
            rows: list[str] = []
            if isinstance(data, list):
                for block in data:
                    if not isinstance(block, dict):
                        continue
                    for _, items in block.items():
                        if not isinstance(items, list):
                            continue
                        for item in items:
                            if not isinstance(item, dict):
                                continue
                            name = str(item.get('name', '')).strip()
                            value = str(item.get('value', '')).strip()
                            if name and value:
                                rows.append(f'{name}: {value}')
            if rows:
                return '\n'.join(rows)
        except Exception:  # noqa: BLE001
            pass

    fallback = re.search(r':short-characteristics=(?P<quote>["\'])(?P<value>(?:.|\n)*?)(?P=quote)', html_text)
    if fallback:
        try:
            data = json.loads(html_lib.unescape(fallback.group('value')))
            rows = []
            if isinstance(data, dict):
                for _, items in data.items():
                    if not isinstance(items, list):
                        continue
                    for item in items:
                        if not isinstance(item, dict):
                            continue
                        name = str(item.get('name', '')).strip()
                        value = str(item.get('value', '')).strip()
                        if name and value:
                            rows.append(f'{name}: {value}')
            if rows:
                return '\n'.join(rows)
        except Exception:  # noqa: BLE001
            pass

    return ""


def extract_description_from_raw_html(html_text: str) -> str:
    raw = _extract_prop_value(html_text, ':descriptions')
    if not raw:
        return ""

    try:
        data = json.loads(raw)
    except Exception:  # noqa: BLE001
        return ""

    blocks: list[str] = []
    if isinstance(data, list):
        for item in data:
            if not isinstance(item, dict):
                continue
            title = str(item.get('title', '')).strip()
            description_html = str(item.get('description', '')).strip()
            if not title and not description_html:
                continue
            description_text = BeautifulSoup(description_html, 'lxml').get_text('\n', strip=True) if description_html else ''
            if title and description_text:
                blocks.append(f'{title}\n{description_text}')
            elif title:
                blocks.append(title)
            elif description_text:
                blocks.append(description_text)

    return '\n\n'.join(blocks).strip()


def extract_photos_from_raw_html(html_text: str) -> list[str]:
    raw = _extract_prop_value(html_text, ':photos')
    if not raw:
        return []

    try:
        data = json.loads(raw)
    except Exception:  # noqa: BLE001
        return []

    photos: list[str] = []
    seen: set[str] = set()
    if isinstance(data, list):
        for item in data:
            if not isinstance(item, dict):
                continue
            for key in ('origin_photo', 'thumb_photo'):
                src = str(item.get(key, '')).strip()
                if not src or src in seen:
                    continue
                seen.add(src)
                if key == 'origin_photo':
                    photos.append(src)
    return photos


def create_openai_client() -> OpenAI:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is not set")
    return OpenAI(api_key=api_key)


def translate_to_montenegrin_latin(
    client: OpenAI,
    product_name: str,
    product_description: str,
    characteristics: str,
    model: str,
) -> dict[str, str]:
    payload = {
        "product_name": product_name,
        "product_description": product_description,
        "characteristics": characteristics,
    }

    response = client.chat.completions.create(
        model=model,
        temperature=0,
        response_format={"type": "json_object"},
        messages=[
            {
                "role": "system",
                "content": (
                    "You translate product catalog text into Montenegrin Latin script. "
                    "Translate from Russian or Ukrainian into natural Montenegrin Latin. "
                    "Preserve product codes, brands, model names, URLs, units, numbers, and line breaks. "
                    "For characteristics, translate labels and keep the value intact unless it clearly needs translation. "
                    "Return only valid JSON with exactly these keys: product_name, product_description, characteristics."
                ),
            },
            {
                "role": "user",
                "content": json.dumps(payload, ensure_ascii=False),
            },
        ],
    )

    content = response.choices[0].message.content or "{}"
    data = json.loads(content)

    return {
        "product_name": str(data.get("product_name", product_name)).strip(),
        "product_description": str(data.get("product_description", product_description)).strip(),
        "characteristics": str(data.get("characteristics", characteristics)).strip(),
    }


def extract_characteristics(section: Tag) -> str:
    table = section.find("table")
    rows: list[str] = []
    if table:
        for tr in table.select("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in tr.find_all(["td", "th"])]
            if len(cells) >= 2:
                rows.append(f"{cells[0]}: {cells[1]}")
    if not rows:
        text = section.get_text("\n", strip=True)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        for idx in range(1, len(lines), 2):
            if idx + 1 < len(lines):
                rows.append(f"{lines[idx]}: {lines[idx + 1]}")
    return "\n".join(rows)


def extract_description(section: Tag) -> str:
    title = section.find(class_=re.compile(r"secondary-title|product-description__title"))
    texts: list[str] = []
    for node in section.find_all(["h2", "h3", "p", "li"]):
        text = node.get_text(" ", strip=True)
        if text:
            texts.append(text)
    if title:
        title_text = title.get_text(" ", strip=True)
        while texts and texts[0] != title_text:
            texts.pop(0)
    cleaned = []
    skip_titles = {"Описание товара", "Описание", "Характеристики", "Показать больше"}
    for text in texts:
        if text in skip_titles:
            continue
        cleaned.append(text)
    return "\n".join(cleaned).strip()


def extract_images(soup: BeautifulSoup) -> list[str]:
    candidates: list[str] = []
    for img in soup.find_all("img"):
        src = img.get("src") or img.get("data-src") or img.get("data-lazy") or img.get("data-original")
        if not src:
            continue
        src = str(src)
        if src.startswith("//"):
            src = "https:" + src
        elif src.startswith("/"):
            src = urljoin(BASE_URL, src)
        if "/cache/products/" not in src:
            continue
        candidates.append(src)

    preferred: list[str] = []
    seen: set[str] = set()
    for src in candidates:
        if src in seen:
            continue
        seen.add(src)
        preferred.append(src)

    originals = [src for src in preferred if "catalog_origin_" in src]
    if len(originals) >= 3:
        return originals[:3]

    mixed = [src for src in preferred if any(token in src for token in ("catalog_origin_", "catalog_product_main_", "order_history_"))]
    if len(mixed) >= 3:
        return mixed[:3]

    return preferred[:3]


def parse_product_page(html: str, source_url: str, requested_code: str) -> ProductRecord:
    soup = BeautifulSoup(html, "lxml")
    main = soup.find("main") or soup

    manufacturer_code = ""
    code_match = re.search(r'\bcode="(\d+)"', html)
    if not code_match:
        code_match = re.search(r'\bproduct-sku="(\d+)"', html)
    if not code_match:
        code_match = re.search(r'\bproductKey&quot;:&quot;(\d+)&quot;', html)
    if code_match:
        manufacturer_code = code_match.group(1)

    product_name = ""
    title_match = re.search(r'\btitle="([^"]+)"', html)
    if title_match:
        product_name = html_lib.unescape(title_match.group(1)).strip()
    if not product_name:
        title_ru_match = re.search(r'\btitle-ru="([^"]+)"', html)
        if title_ru_match:
            product_name = html_lib.unescape(title_ru_match.group(1)).strip()
    if not product_name:
        h1 = main.find("h1")
        if h1:
            product_name = h1.get_text(" ", strip=True)
    if not product_name:
        page_title = soup.find("title")
        if page_title:
            product_name = page_title.get_text(" ", strip=True)

    description = extract_description_from_raw_html(html)
    characteristics = extract_characteristics_from_raw_html(html)
    images = extract_photos_from_raw_html(html)

    if not description or not characteristics or not images:
        fallback_description_section = soup.find(id="description")
        fallback_characteristics_section = soup.find(id="characteristics")
        if not description and fallback_description_section:
            description = extract_description(fallback_description_section)
        if not characteristics and fallback_characteristics_section:
            characteristics = extract_characteristics(fallback_characteristics_section)
        if not images:
            images = extract_images(soup)

    return ProductRecord(
        article_input=requested_code,
        manufacturer_code=manufacturer_code,
        product_name=product_name,
        product_description=description,
        characteristics=characteristics,
        image_1=images[0] if len(images) > 0 else "",
        image_2=images[1] if len(images) > 1 else "",
        image_3=images[2] if len(images) > 2 else "",
        source_url=source_url,
    )


def translate_record(
    record: ProductRecord,
    client: OpenAI,
    cache: dict[str, dict[str, str]],
    model: str,
) -> ProductRecord:
    cache_key = make_translation_key(record.product_name, record.product_description, record.characteristics)
    cached = cache.get(cache_key)
    if cached:
        record.product_name = cached.get("product_name", record.product_name)
        record.product_description = cached.get("product_description", record.product_description)
        record.characteristics = cached.get("characteristics", record.characteristics)
        return record

    translated = translate_to_montenegrin_latin(
        client=client,
        product_name=record.product_name,
        product_description=record.product_description,
        characteristics=record.characteristics,
        model=model,
    )
    record.product_name = translated["product_name"]
    record.product_description = translated["product_description"]
    record.characteristics = translated["characteristics"]
    cache[cache_key] = translated
    return record


def build_article_index(
    session: requests.Session,
    product_urls: list[str],
    requested_codes: list[str],
    reindex: bool = False,
    workers: int = 16,
) -> dict[str, str]:
    cached = load_article_index()
    normalized_requested = [canonical_lookup_code(code) for code in requested_codes if canonical_lookup_code(code)]
    normalized_cached = {normalize_code(key): value for key, value in cached.items() if normalize_code(key)}
    if not reindex:
        return normalized_cached
    if normalized_cached and all(code in normalized_cached for code in normalized_requested):
        return normalized_cached

    index: dict[str, str] = {}

    def probe(url: str) -> tuple[str, str] | None:
        try:
            html = fetch_html_with_recovery(session, url)
        except Exception:
            return None

        extracted_code = extract_product_code_from_html(html)
        if not extracted_code:
            return None

        return extracted_code, url

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(probe, url): url for url in product_urls}
        for position, future in enumerate(as_completed(futures), start=1):
            result = future.result()
            if result:
                code, url = result
                index[code] = url
            if position % 200 == 0:
                print(f"Indexed {position}/{len(product_urls)} product pages", file=sys.stderr)

    save_article_index(index)
    return index


def collect_product(
    session: requests.Session,
    article_index: dict[str, str],
    product_urls: list[str],
    code: str,
) -> ProductRecord:
    normalized = canonical_lookup_code(code)
    url = article_index.get(normalized, "")
    if not url:
        discovered = discover_product_url(session, normalized, product_urls)
        if discovered:
            article_index[normalized] = discovered
            save_article_index(article_index)
            url = discovered
            log_event("info", "url_discovered_from_fallback", code=normalized, url=discovered)
        else:
            log_event("warn", "article_not_found_after_fallback", code=normalized)
            return ProductRecord(article_input=code, status="not_found", error="Товар не найден в sitemap и fallback-поиск")

    try:
        html = fetch_html_with_recovery(session, url)
    except Exception as exc:  # noqa: BLE001
        log_event("warn", "fetch_failed_primary_url", code=normalized, url=url, error=str(exc))
        discovered = discover_product_url(session, normalized, product_urls)
        if discovered and discovered != url:
            try:
                html = fetch_html_with_recovery(session, discovered)
                url = discovered
                article_index[normalized] = discovered
                save_article_index(article_index)
                log_event("info", "fetch_recovered_with_discovered_url", code=normalized, url=discovered)
            except Exception as discovered_exc:  # noqa: BLE001
                log_event(
                    "warn",
                    "discovered_url_fetch_failed",
                    code=normalized,
                    url=discovered,
                    error=str(discovered_exc),
                )
                html = ""
        else:
            html = ""

        fallback_candidates = [
            url.replace("/ru/", "/ua/"),
            url.replace("/ua/", "/ru/"),
            url.replace(f"{BASE_URL}/ru/", f"{BASE_URL}/"),
            url.replace(f"{BASE_URL}/ua/", f"{BASE_URL}/"),
        ]
        deduped_fallbacks = list(dict.fromkeys([candidate for candidate in fallback_candidates if candidate]))
        if not html:
            for fallback_url in deduped_fallbacks:
                try:
                    html = fetch_html_with_recovery(session, fallback_url)
                    url = fallback_url
                    article_index[normalized] = fallback_url
                    save_article_index(article_index)
                    log_event("info", "fetch_recovered_with_alt_url", code=normalized, url=fallback_url)
                    break
                except Exception as fallback_exc:  # noqa: BLE001
                    log_event(
                        "warn",
                        "alt_url_fetch_failed",
                        code=normalized,
                        url=fallback_url,
                        error=str(fallback_exc),
                    )
                    continue

        if not html:
            discovered = discover_product_url(session, normalized, product_urls)
            if discovered and discovered != url:
                try:
                    html = fetch_html_with_recovery(session, discovered)
                    url = discovered
                    article_index[normalized] = discovered
                    save_article_index(article_index)
                    log_event("info", "fetch_recovered_with_discovered_url_second_pass", code=normalized, url=discovered)
                except Exception:
                    html = ""

        if not html:
            log_event("error", "fetch_error_after_all_recovery", code=normalized, url=url, error=str(exc))
            return ProductRecord(
                article_input=code,
                manufacturer_code=normalized,
                source_url=url,
                status="fetch_error",
                error=str(exc),
            )

    record = parse_product_page(html, url, code)
    if not record.manufacturer_code:
        record.manufacturer_code = normalized
    if record.product_name:
        record.status = "ok"
        record.error = ""
    elif not record.product_name:
        record.status = "parse_error"
        record.error = "Не удалось извлечь название товара"
    return record


def write_excel(records: list[ProductRecord], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"

    headers = [
        "article_input",
        "Šifra proizvođača",
        "Naziv proizvoda",
        "Opis proizvoda",
        "Karakteristike",
        "Slike",
        "source_url",
        "status",
        "error",
    ]
    sheet.append(headers)

    for record in records:
        sheet.append([
            record.article_input,
            record.manufacturer_code,
            record.product_name,
            record.product_description,
            record.characteristics,
            "\n".join([value for value in [record.image_1, record.image_2, record.image_3] if value]),
            record.source_url,
            record.status,
            record.error,
        ])

    for row in sheet.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column in sheet.columns:
        max_len = 0
        column_letter = column[0].column_letter
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        sheet.column_dimensions[column_letter].width = min(max_len + 2, 90)

    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Dnipro-M scraper to Excel")
    parser.add_argument("--input", default="articles.txt", help="Path to a file with article codes (.txt, .csv, .xlsx)")
    parser.add_argument("--articles", nargs="*", default=[], help="Article codes passed directly on the command line")
    parser.add_argument("--output", default="dnipro_m_products.xlsx", help="Excel output path")
    parser.add_argument("--max-workers", type=int, default=10, help="Concurrent page fetches while searching")
    parser.add_argument("--translate-model", default=os.getenv("OPENAI_MODEL", "gpt-4o-mini"), help="OpenAI model for Montenegrin Latin translation")
    parser.add_argument("--no-translate", action="store_true", help="Skip OpenAI translation")
    parser.add_argument("--cache-only", action="store_true", help="Deprecated alias for the default cache-only mode")
    parser.add_argument("--reindex", action="store_true", help="Rebuild the article URL index from sitemap before processing")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_path = Path(args.output)

    input_path = Path(args.input) if args.input else None
    if args.articles and args.input == "articles.txt":
        input_path = None

    if input_path is not None and not input_path.exists():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        return 1

    codes = load_article_codes(input_path, args.articles)

    if not codes:
        if input_path is not None and input_path.name == "articles.txt":
            print("articles.txt is empty; add article codes or pass --articles", file=sys.stderr)
        else:
            print("No article codes found in input file", file=sys.stderr)
        return 1

    ensure_cache_dir()
    session = requests.Session()

    product_urls = get_product_urls(session)
    if not product_urls:
        print("No product URLs found in sitemap", file=sys.stderr)
        return 1

    if args.cache_only and args.reindex:
        print("Use either --cache-only or --reindex, not both.", file=sys.stderr)
        return 1

    article_index = build_article_index(session, product_urls, codes, reindex=args.reindex, workers=args.max_workers)
    if not article_index:
        if not args.reindex:
            print("Cache-only mode is on, but the article URL cache is empty or missing.", file=sys.stderr)
        else:
            print("No article codes could be indexed from product pages", file=sys.stderr)
        return 1

    translation_client: OpenAI | None = None
    translation_cache: dict[str, dict[str, str]] = {}
    if not args.no_translate:
        try:
            translation_client = create_openai_client()
            translation_cache = load_translation_cache()
        except Exception as exc:  # noqa: BLE001
            print(f"OpenAI translation is enabled, but cannot start: {exc}", file=sys.stderr)
            print("Set OPENAI_API_KEY or run with --no-translate for a local-only export.", file=sys.stderr)
            return 1

    start = time.time()
    records: list[ProductRecord] = []
    for index, code in enumerate(codes, start=1):
        print(f"[{index}/{len(codes)}] Processing {code}...")
        log_event("info", "processing_started", article=code, index=index, total=len(codes))
        record = collect_product(session, article_index, product_urls, code)
        if record.status == "ok" and translation_client is not None:
            try:
                record = translate_record(record, translation_client, translation_cache, args.translate_model)
            except Exception as exc:  # noqa: BLE001
                record.status = "translation_error"
                record.error = str(exc)
        records.append(record)
        log_event(
            "info",
            "processing_finished",
            article=code,
            status=record.status,
            source_url=record.source_url,
            error=record.error,
        )
        if record.status != "ok":
            print(f"  -> {record.status}: {record.error}")

    if translation_client is not None:
        save_translation_cache(translation_cache)

    write_excel(records, output_path)
    elapsed = time.time() - start
    print(f"Saved {len(records)} rows to {output_path} in {elapsed:.1f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())